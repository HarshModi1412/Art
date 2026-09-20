#!/usr/bin/env python3
"""Builds the One Tap Manager Instagram prospecting kit (xlsx)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INK = "1A1C21"; SLATE = "5C6790"; SOFT = "EEF0F6"; LINE = "D7DAE1"; MUT = "6B7280"
HEADFILL = PatternFill("solid", fgColor=SLATE)
SOFTFILL = PatternFill("solid", fgColor=SOFT)
thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE = Font(name="Calibri", bold=True, color=INK, size=15)
SUB = Font(name="Calibri", color=MUT, size=10, italic=True)
B = Font(name="Calibri", bold=True, color=INK, size=11)
N = Font(name="Calibri", color=INK, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- Sheet 1: read me / targeting
ws = wb.active; ws.title = "Targeting"
ws.sheet_view.showGridLines = False
ws["A1"] = "One Tap Manager — Instagram prospecting"; ws["A1"].font = TITLE
ws["A2"] = ("Who to approach, where to find them, and how to qualify. Built from Brand.md. "
            "Fill the Tracker sheet as you (or the script) collect accounts.")
ws["A2"].font = SUB; ws.merge_cells("A2:F2")

def section(ws, r, title):
    ws.cell(r, 1, title).font = B
    ws.cell(r, 1).fill = SOFTFILL
    for c in range(1, 7):
        ws.cell(r, c).fill = SOFTFILL; ws.cell(r, c).border = BORDER
    return r + 1

def rows(ws, r, pairs, wcol=1):
    for k, v in pairs:
        ws.cell(r, 1, k).font = B; ws.cell(r, 1).alignment = TOP
        c = ws.cell(r, 2, v); c.font = N; c.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1
    return r + 1

r = 4
r = section(ws, r, "Who is a fit (Ideal Customer Profile)")
r = rows(ws, r, [
 ("In one line", "A small Indian D2C seller running the shop themselves (or with 1–2 helpers), already making sales on Instagram/WhatsApp, short on time, tired of paying for a stack of apps."),
 ("Core niches", "Women's ethnic & fusion wear (saree, kurti, lehenga, suits); handmade / oxidised / imitation jewellery; attar & perfume; small-batch beauty & skincare; candles & home; food (pickles, snacks, bakery, chocolate)."),
 ("Sweet-spot size", "~1,000–50,000 followers. Real sales, not a big brand with a team. Skip 100k+ polished brands (they have staff and a full site) and skip <500 (usually dead or hobby)."),
 ("Strong buy signals", "Bio or posts say “DM to order”, “WhatsApp to order”, price in comments, COD, “ships all over India”; a wa.me / Linktree link rather than a real store; reseller-style catalogue posts."),
 ("NOT a fit", "Enterprises, multi-outlet chains, manufacturers, anyone with a finance team, and slick brands already on full Shopify with an agency. These are section 2.7 of Brand.md — not for us."),
])
r = section(ws, r, "Two qualification gates (both must pass)")
r = rows(ws, r, [
 ("1. Alive", "Last post is 30 days old or newer. Dead accounts waste outreach. The script checks the newest post date; by hand, look at the top-left post."),
 ("2. The website angle", "Record whether they have NO site, a Linktree/social-only link, a basic site, or a full store (Shopify/Dukaan/Instamojo). Note it either way — the pitch differs (see Outreach)."),
])
r = section(ws, r, "Where to search on Instagram (seed the list from these)")
r = rows(ws, r, [
 ("Order-signal tags", "#dmtoorder  #whatsapptoorder  #shoponinstagram  #instashopindia  #codavailable  #shippingalloverindia  #onlineshoppingindia  #supportsmallbusiness  #smallbusinessindia  #madeinindia"),
 ("Clothing", "#boutiqueindia  #kurti  #sareelove  #lehenga  #suitsonline  #ethnicwear  #fusionwear  #cottonkurti  #handblockprint"),
 ("Jewellery", "#oxidisedjewellery  #imitationjewellery  #handmadejewellery  #jewelleryindia  #antiquejewellery  #kundanjewellery"),
 ("Perfume / beauty / food", "#attar  #perfumeindia  #fragranceindia  #skincareindia  #handmadesoap  #candlesindia  #homemadefood  #bakerylife  #pickleslove"),
 ("City tags (high-density seller cities)", "#suratboutique  #jaipurjewellery  #ahmedabadboutique  #delhiwholesale  #chandnichowk  #mumbaifashion  #bangaloreshopping  #hyderabadshopping  #lucknowchikankari  #kolkatafashion  #indoreboutique  #ludhianasuits"),
 ("How to expand", "Open one good fit, then work its Followers/Following and Instagram's “Suggested for you” — sellers in the same niche cluster together. This is how you get from 10 to 100."),
])
r = section(ws, r, "How to use this kit")
r = rows(ws, r, [
 ("Fastest path (verified)", "Run collect_instagram_prospects.py where Instagram is reachable, with an Apify token (recommended) or your logged-in IG session. It fills the Tracker with real handles, last-post dates and website flags."),
 ("By hand", "Search a tag above, open accounts, apply the two gates, and paste rows into the Tracker. ~30–40 qualified/hour once you have the eye for it."),
 ("Rule", "Never add an account you have not actually opened. A list of guesses is worse than a short real one."),
])
ws.column_dimensions["A"].width = 30
for col in "BCDEF": ws.column_dimensions[col].width = 22

# ---------------------------------------------------------------- Sheet 2: tracker
tk = wb.create_sheet("Tracker")
tk.sheet_view.showGridLines = False
cols = ["#", "Handle (@)", "Profile URL", "Niche", "City / region", "Followers",
        "Last post date", "Active (≤30d)", "Has website?", "Website type",
        "Website URL", "Order method", "Fit 1-3", "Why a fit", "Outreach note",
        "Status", "Date contacted"]
widths = [4, 20, 30, 20, 16, 11, 14, 12, 12, 18, 26, 16, 8, 30, 34, 14, 14]
tk["A1"] = "Prospect tracker"; tk["A1"].font = TITLE
tk["A2"] = ("One row per real, opened account. Columns 7-10 are the two gates. "
            "'Has website?' = No / Linktree / Basic / Full store. Fit: 3 = solo seller, active, DM/WhatsApp-led, no real site.")
tk["A2"].font = SUB; tk.merge_cells("A2:Q2")
hr = 4
for i, (name, w) in enumerate(zip(cols, widths), 1):
    c = tk.cell(hr, i, name); c.font = H; c.fill = HEADFILL; c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    tk.column_dimensions[get_column_letter(i)].width = w
tk.row_dimensions[hr].height = 30
# one worked example row so the format is obvious, clearly marked as a sample
example = [1, "(example) @surat_kurti_studio", "https://instagram.com/…", "Cotton kurtis",
           "Surat", "8,400", "18 Sep 2026", "Yes", "Linktree", "Linktree",
           "linktr.ee/…", "DM + WhatsApp", 3, "Solo seller, posts daily, orders via DM, no real store",
           "Lead with free analytics + own website with no badge", "To contact", ""]
for i, v in enumerate(example, 1):
    c = tk.cell(hr + 1, i, v); c.font = Font(italic=True, color=MUT, size=10)
    c.alignment = WRAP; c.border = BORDER
for rr in range(hr + 2, hr + 2 + 100):
    tk.cell(rr, 1, rr - hr - 1)
    for i in range(1, len(cols) + 1):
        cell = tk.cell(rr, i); cell.border = BORDER; cell.alignment = WRAP; cell.font = N
tk.freeze_panes = "A5"

# ---------------------------------------------------------------- Sheet 3: outreach + method
ot = wb.create_sheet("Outreach & method")
ot.sheet_view.showGridLines = False
ot["A1"] = "Outreach & method"; ot["A1"].font = TITLE
r = 3
r = section(ot, r, "DM — no proper website (the strongest fit)")
r = rows(ot, r, [
 ("", "Hi <name>, love what you're doing with <product>. Quick one: right now your orders come through DMs. "
      "One Tap Manager gives you your own selling website with no badge, reads the sales you already make, and "
      "tells you the two or three things worth doing each morning. Analytics is free forever, no cut of your sales, "
      "flat ₹999 for the rest. Want me to set your site up so you can see it?"),
])
r = section(ot, r, "DM — already has a basic site / Linktree")
r = rows(ot, r, [
 ("", "Hi <name>, your <product> line looks great. You're already selling online, so the gap is usually the back office: "
      "who reorders, who to win back, a week of posts written and scheduled. One Tap Manager does that work, not just the "
      "charts, and replaces four to six apps with one flat ₹999, never a cut of your sales. Happy to show you on your own numbers, "
      "takes about two minutes to connect."),
])
r = section(ot, r, "Voice rules (from Brand.md — keep these)")
r = rows(ot, r, [
 ("Do", "Plain words, one idea at a time. Warm and India-first. Rupees, WhatsApp, COD, Hinglish where natural. Say the price and that we never take a cut."),
 ("Don't", "No em dashes. No emoji as icons. No hype or fake claims ('40% more sales'). No fake scarcity. Sentence case."),
])
r = section(ot, r, "Building the verified list")
r = rows(ot, r, [
 ("Why not from the sandbox", "This cloud environment is network-blocked from Instagram and IG needs a logged-in session for profile/post data, so the list is built where IG is reachable, with your session or a scraping service."),
 ("Recommended", "Apify actors (apify/instagram-hashtag-scraper then apify/instagram-profile-scraper) — they handle anti-bot and return post timestamps + bio external URL. Free tier covers a first run. Put the token in the script."),
 ("Fallback", "instaloader with your logged-in session file — works but rate-limits fast; go slow (a few hundred profiles/day) to keep the account safe."),
 ("Compliance note", "Automated scraping is against Instagram's Terms. Keep volumes low, respect rate limits, target business accounts only, and prefer a service that handles this. This is a lead-research aid, not a bulk-harvest tool."),
])
ot.column_dimensions["A"].width = 26
for col in "BCDEF": ot.column_dimensions[col].width = 24

wb.save("/mnt/user-data/outputs/otm-instagram-prospecting.xlsx")
print("wrote otm-instagram-prospecting.xlsx")
